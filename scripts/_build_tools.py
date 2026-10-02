# -*- coding: utf-8 -*-
"""実家じまいチェックリスト・業者見積もり比較シートを生成する（HTML + xlsx）。
実行: PYTHONUTF8=1 python _build_tools.py
出力: jikkajimai-checklist.html / jikkajimai-checklist.xlsx / mitsumori-hikaku.html / mitsumori-hikaku.xlsx
"""
import html
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
SITE = "にいがた実家じまい帖"
UPDATED = "2026-10-03"
ROBOTS = '<meta name="robots" content="noindex,nofollow,noarchive,noimageindex">'
e = html.escape

# ------------------------------------------------------------------ 共通CSS
BASE_CSS = """
:root{--ink:#37352f;--sub:#6b6a66;--line:#d9d8d4;--soft:#f7f7f5}
*{box-sizing:border-box}
html{-webkit-text-size-adjust:100%}
body{margin:0;background:#fff;color:var(--ink);font-family:"Hiragino Sans","Hiragino Kaku Gothic ProN","Yu Gothic UI","Yu Gothic","Meiryo",sans-serif;line-height:1.6;font-size:15px}
.wrap{max-width:900px;margin:0 auto;padding:24px 16px 48px}
h1{font-size:22px;margin:0 0 4px;font-weight:700}
h2{font-size:16px;margin:28px 0 8px;padding-bottom:4px;border-bottom:1px solid var(--ink)}
.meta{color:var(--sub);font-size:13px;margin:0 0 12px}
.note{font-size:13px;color:var(--sub)}
.bar{display:flex;gap:8px;align-items:center;margin:0 0 16px;flex-wrap:wrap}
.btn{font:inherit;font-size:14px;padding:6px 14px;border:1px solid var(--ink);background:#fff;color:var(--ink);border-radius:4px;cursor:pointer}
table{border-collapse:collapse;width:100%}
th,td{border:1px solid var(--line);padding:6px 8px;vertical-align:top;text-align:left}
th{background:var(--soft);font-weight:700;font-size:13px}
.footer{margin-top:24px;padding-top:8px;border-top:1px solid var(--line);font-size:12px;color:var(--sub)}
input[type=checkbox]{appearance:none;-webkit-appearance:none;width:15px;height:15px;border:1.5px solid #000;border-radius:2px;margin:0;vertical-align:-2px;background:#fff;cursor:pointer;position:relative}
input[type=checkbox]:checked::after{content:"";position:absolute;left:3px;top:-1px;width:4px;height:9px;border:solid #000;border-width:0 2px 2px 0;transform:rotate(45deg)}
"""

# ------------------------------------------------------------------ 1. チェックリスト
# (区分, [(項目, 期限・目安, 確認先), ...])
CHECKLIST = [
    ("1. 死後・施設入所後の期限つき手続き", [
        ("年金受給権者の死亡届を出す", "すみやかに。遅れると受け取りすぎた年金を返す場合がある", "日本年金機構、市の公式ページ（死亡後の手続き）"),
        ("相続放棄をするか決める（決めるまで遺品の処分・売却を急がない）", "相続の開始を知ったときから3か月以内", "家庭裁判所、市の公式ページ（法律相談）"),
        ("準確定申告が必要か確認する", "相続の開始を知った日の翌日から4か月以内", "税務署、市の公式ページ（税の相談）"),
        ("相続税の申告・納税が必要か確認する", "死亡を知った日の翌日から10か月以内", "税務署、市の公式ページ（税の相談）"),
        ("相続登記（不動産の名義変更）の準備をする", "取得を知った日から3年以内（2024年4月1日から義務化）", "法務局、市の公式ページ（相続登記の案内）"),
        ("空き家の3,000万円特別控除が使えるか確認する", "相続の開始から3年を経過する日の属する年の12月31日まで。令和9年12月31日までの売却が条件", "税務署、市の公式ページ（空き家の譲渡）"),
        ("施設入所の場合：家の鍵・郵便・火の元・公共料金の管理役を決める", "入所が決まったらすぐ", "入所先の施設、市の公式ページ（高齢者福祉）"),
    ]),
    ("2. 貴重品・書類の確認", [
        ("通帳・印鑑・キャッシュカードを探して別にしておく", "ごみ袋に入れる前に", "各金融機関、市の公式ページ（相続の手続き）"),
        ("保険証券・年金関係の書類を探す", "ごみ袋に入れる前に", "各保険会社、日本年金機構"),
        ("権利証（登記関係の書類）・固定資産税の通知書を探す", "ごみ袋に入れる前に", "市の公式ページ（固定資産税）、法務局"),
        ("契約書（賃貸・ローン・保険など）、鍵、暗証番号のメモを探す", "ごみ袋に入れる前に", "各契約先"),
        ("現金・貴金属を探す（タンスの引き出し、仏壇まわり、本の間も）", "ごみ袋に入れる前に", "市の公式ページ（相続の手続き）"),
        ("見つけた物は中身を確認せず、一つの箱にまとめる", "仕分けの途中で", "市の公式ページ（相続の手続き）"),
    ]),
    ("3. ごみの出し方（市の粗大ごみ・持ち込み・家電）", [
        ("市の粗大ごみの申込方法を調べる（事前申込か、持ち込みのみか）", "収集まで数日〜2週間ほどかかる市もある。早めに", "市の公式ページ（粗大ごみ）。30市町村の一覧は「粗大ごみ早見表」にもあります"),
        ("品目ごとの手数料（処理券）の額と、買える場所を確認する", "申込時に金額を案内される市が多い", "市の公式ページ（粗大ごみ手数料）"),
        ("持ち込み先の受付時間・休み・予約の要否を確認する", "日曜・祝日・年末年始は休みの施設が多い", "市の公式ページ（ごみ処理施設）"),
        ("混む日を避けて持ち込む日を決める", "新潟市は年末年始・お盆・引っ越しシーズン・連休の翌日・土曜が混むと案内", "市の公式ページ（自己搬入）"),
        ("家電4品目（エアコン・テレビ・冷蔵庫/冷凍庫・洗濯機/衣類乾燥機）は市で出せないと確認する", "販売店か指定引取場所へ。リサイクル料金が別にかかる", "市の公式ページ（家電リサイクル）"),
        ("納屋・車庫の「市で出せない物」を分ける（農薬、バッテリー、消火器、ガスボンベ、灯油、塗料、バイク、タイヤなど）", "市ごとに対象が違う", "市の公式ページ（市で収集しないごみ）"),
        ("袋・結び方・大きさなど、市の出し方の条件を確認する", "条件を外れると収集されない場合がある", "市の公式ページ（ごみの出し方）"),
    ]),
    ("4. 業者に頼む範囲", [
        ("頼む範囲を決める（搬出と処分だけ／仕分けから／清掃・買取まで）", "量・階段・雪・遠方・期限の5つで考える", "市の公式ページ（許可業者の案内）"),
        ("家庭ごみを運べる許可（一般廃棄物の収集運搬）の有無を、市の一覧で確認する", "産業廃棄物や古物商の許可では家庭の廃棄物は回収できない", "市の公式ページ（ごみ収集運搬業許可業者）"),
        ("実家のある地域に対応している業者か確認する", "地域限定の許可がある市もある", "市の公式ページ（許可業者一覧）"),
        ("複数社から見積もりを取り、書面で受け取る（別紙の比較シートを使う）", "追加料金・キャンセル料・作業範囲を書面で確認", "市の公式ページ、消費生活センター"),
        ("売れそうな物は、先に買取を確認する", "処分する量が減る", "市の公式ページ（リサイクル・再利用）"),
        ("搬出・処分費の補助がないか確認する", "契約前の申請が必要な制度がある", "市の公式ページ（空き家・補助金）"),
    ]),
    ("5. 空き家の方針", [
        ("建物をどうするか方針を決める（売る・貸す・使う・壊す）", "すぐ決められなくてもよい。放置は避ける", "市の公式ページ（空き家対策）"),
        ("空き家バンクの登録条件を確認する", "市によって有無・条件が違う", "市の公式ページ（空き家バンク）"),
        ("解体の補助金の有無・条件を確認する", "着工前の申請が必要な制度がある", "市の公式ページ（空き家解体補助）"),
        ("管理不全空家として指導を受けない管理をする（見回り・換気・通水）", "指導に従わず勧告を受けると、固定資産税の住宅用地特例が受けられなくなる", "市の公式ページ（空き家の管理）"),
        ("相談窓口の電話番号を控える", "長岡市は都市政策課、新潟市は住環境政策課など", "市の公式ページ（空き家相談）"),
    ]),
    ("6. 名義・公共料金・郵便", [
        ("電気・ガス・水道を解約するか名義変更するか決め、各社へ連絡する", "解約日は片付けの日程に合わせる", "各事業者、市の公式ページ（上下水道）"),
        ("長期不在になる場合の止水栓の扱いを確認する", "冬の凍結にも関わる", "市の公式ページ（水道）"),
        ("郵便物の転送手続きをする（転居届は届出日から1年間、無料で転送）", "亡くなった方あての郵便は窓口で確認", "日本郵便、市の公式ページ"),
        ("戸籍の取り寄せ（本籍地以外の市区町村の窓口でも請求できる）", "窓口に本人が出向く必要がある", "市の公式ページ（戸籍証明書の広域交付）"),
        ("固定資産税の納税通知書の送付先と、納める人を確認する", "納期を確認", "市の公式ページ（固定資産税）"),
    ]),
    ("7. 近所への挨拶", [
        ("片付けの日程（トラックが入る日・音が出る日）を近所に伝える", "できれば片付けを始める前に", "市の公式ページ（自治会・町内会）"),
        ("自分の連絡先を、隣家や自治会・町内会に渡す", "空き家になる場合は特に", "市の公式ページ（空き家の管理）"),
    ]),
    ("8. 冬の雪対策", [
        ("雪が降る前に搬出まで終える日程にする", "雪の時期は搬出路が滑りやすい", "市の公式ページ（雪対策）"),
        ("空き家の雪下ろし・除雪の責任と、誰がやるかを決める", "新潟県は「空き家の管理は所有者の責任です」と案内", "市の公式ページ（雪下ろし・除雪の支援）"),
        ("雪下ろしの支援・補助制度の有無を確認する", "市ごとに制度が違う", "市の公式ページ（雪下ろし支援）"),
        ("業者の積雪期の対応（日程・料金の条件が変わるか）を見積もりで聞く", "見積もりの段階で確認", "市の公式ページ、業者"),
        ("雪の中の作業は事故に注意する（雪下ろしだけでなく作業全般）", "新潟県が除雪作業事故の防止を呼びかけている", "新潟県・市の公式ページ（除雪事故防止）"),
    ]),
]
CHECK_NOTE = ("期限と手続きの要否は、一般的な案内です。自分に当てはまるかは、税理士・司法書士・弁護士・税務署などで確認してください。"
              "市の制度は更新されます。申込の前に、市の公式ページで最新の内容を確認してください。")

CSS_CHECK = BASE_CSS + """
.cat{margin-top:18px}
table.cl td.ck{width:28px;text-align:center}
table.cl td.dl{width:26%;font-size:12.5px}
table.cl td.cf{width:25%;font-size:12.5px}
.who{display:flex;gap:16px;flex-wrap:wrap;margin:8px 0 4px;font-size:13px}
.who span{border-bottom:1px solid var(--ink);min-width:150px;display:inline-block;height:1.4em}
@page{size:A4 portrait;margin:12mm}
@media print{
  body{font-size:10pt;line-height:1.45;color:#000}
  .wrap{max-width:none;padding:0}
  .bar{display:none}
  h1{font-size:16pt}
  h2{font-size:11pt;margin:12px 0 4px;border-bottom:1.5px solid #000;page-break-after:avoid}
  th,td{border-color:#000;padding:3px 5px}
  th{background:#fff;border-bottom:1.5px solid #000}
  tr{page-break-inside:avoid}
  .meta,.note,.footer{color:#000}
}
"""


def build_checklist_html():
    rows = []
    for cat, items in CHECKLIST:
        body = "".join(
            f'<tr><td class="ck"><input type="checkbox" aria-label="済"></td><td>{e(t)}</td><td class="dl">{e(d)}</td><td class="cf">確認先: {e(c)}</td></tr>'
            for t, d, c in items
        )
        rows.append(
            f'<section class="cat"><h2>{e(cat)}</h2><table class="cl"><thead><tr><th>済</th><th>確認項目</th><th>期限・目安</th><th>確認先</th></tr></thead><tbody>{body}</tbody></table></section>'
        )
    doc = f"""<!doctype html>
<html lang="ja">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
{ROBOTS}
<title>実家じまいチェックリスト（新潟版）</title>
<style>{CSS_CHECK}</style>
</head>
<body>
<div class="wrap">
<h1>実家じまいチェックリスト（新潟版）</h1>
<p class="meta">{SITE}／更新 {UPDATED}／A4縦・白黒印刷用／登録不要・無料・再配布可（出典明記）</p>
<div class="bar"><button class="btn" onclick="window.print()">印刷する</button><span class="note">チェックはこの画面でも付けられます（保存はされません）。</span></div>
<div class="who"><div>実家の所在地（市町村）: <span></span></div><div>記入日: <span></span></div><div>担当者: <span></span></div></div>
{''.join(rows)}
<p class="footer">{e(CHECK_NOTE)}<br>出典: {SITE}「新潟の実家じまいの手順」「新潟の実家の片付けは自分で？業者？」（各機関の公式情報に基づく）。このシートは出典を明記すれば自由に再配布できます。</p>
</div>
</body>
</html>
"""
    with open(os.path.join(HERE, "jikkajimai-checklist.html"), "w", encoding="utf-8", newline="\n") as f:
        f.write(doc)


# ------------------------------------------------------------------ xlsx 共通スタイル
FONT = "Yu Gothic"
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill("solid", fgColor="F1F1EF")
HINT_FILL = PatternFill("solid", fgColor="FAFAF8")


def style_cell(c, bold=False, fill=None, wrap=True, size=10, color="37352F", align="left", valign="top"):
    c.font = Font(name=FONT, size=size, bold=bold, color=color)
    c.alignment = Alignment(wrap_text=wrap, vertical=valign, horizontal=align)
    c.border = BORDER
    if fill:
        c.fill = fill


def build_checklist_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "チェックリスト"
    ws["A1"] = "実家じまいチェックリスト（新潟版）"
    ws["A1"].font = Font(name=FONT, size=14, bold=True)
    ws["A2"] = f"{SITE}／更新 {UPDATED}／登録不要・無料・再配布可（出典明記）"
    ws["A2"].font = Font(name=FONT, size=9, color="6B6A66")
    ws["A3"] = "実家の所在地（市町村）:"
    ws["A3"].font = Font(name=FONT, size=10)
    ws["C3"] = "記入日:"
    ws["C3"].font = Font(name=FONT, size=10)
    heads = ["区分", "確認項目", "期限・目安", "確認先", "済", "メモ"]
    r = 5
    for i, h in enumerate(heads, 1):
        style_cell(ws.cell(r, i, h), bold=True, fill=HEAD_FILL, valign="center")
    r += 1
    for cat, items in CHECKLIST:
        for t, d, c in items:
            vals = [cat.split(". ", 1)[1], t, d, "確認先: " + c, "□", ""]
            for i, v in enumerate(vals, 1):
                cell = ws.cell(r, i, v)
                style_cell(cell, size=10, align="center" if i == 5 else "left", valign="center" if i == 5 else "top")
            ws.cell(r, 5).font = Font(name=FONT, size=12)
            r += 1
    r += 1
    ws.cell(r, 1, CHECK_NOTE).font = Font(name=FONT, size=9, color="6B6A66")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 30
    for col, w in zip("ABCDEF", [18, 46, 34, 34, 5, 24]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A6"
    ws.print_title_rows = "5:5"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    wb.save(os.path.join(HERE, "jikkajimai-checklist.xlsx"))


# ------------------------------------------------------------------ 2. 見積もり比較シート
# (項目, 確認の目安, 入力種別)  入力種別: text / yen / list:<options> / total
MITSUMORI_ROWS = [
    ("業者名", "", "name"),
    ("許可の有無（一般廃棄物収集運搬）", "市の公式一覧で確認する。運搬を行う会社が許可業者かも尋ねる。産業廃棄物・古物商の許可では家庭の廃棄物は回収できない", "list:あり,なし,未確認"),
    ("見積もり方法（訪問・写真）", "訪問見積もりか、写真見積もりか。写真の場合、当日の変更の扱いを聞く", "list:訪問,写真,両方"),
    ("基本料金（円）", "何が含まれるか（作業人数・車両・仕分け・清掃）を書面で確認", "yen"),
    ("追加料金の条件", "階段・駐車場・量の増加・特殊な品目で増えるか。キャンセル料はいつから・いくらか", "text"),
    ("追加料金の見込み額（円）", "条件に当てはまる場合の見込み額。分からなければ空欄", "yen"),
    ("買取の有無・金額", "売れる物（家具・着物・家電など）の買取があるか。処分費と相殺されるか", "text"),
    ("買取金額（円）", "合計から差し引く金額", "yen"),
    ("作業日・立会い", "作業日の候補。立会いの要否、鍵の預け方と返却", "text"),
    ("積雪期の対応", "雪の時期に日程・料金の条件が変わるか", "text"),
    ("写真報告の有無", "作業前後の写真報告があるか（遠方の場合は特に）", "list:あり,なし"),
    ("合計（円）", "基本料金 + 追加料金の見込み額 - 買取金額（自動計算）", "total"),
    ("備考", "", "text"),
]
SIX = [
    ("一般廃棄物の収集運搬の許可はありますか。運ぶのは自社ですか", "市の公式一覧にある業者か確認する。遺品整理の会社が別の会社に運搬を任せる場合、その運搬会社が許可業者か確認する"),
    ("追加料金は、どんな場合に、いくら増えますか", "階段・駐車場・量の増加・特殊な品目。見積書に追加料金の条件を書いてもらう"),
    ("キャンセル料は、いつから、いくらかかりますか", "日程変更の扱いも聞く。後から高額なキャンセル料を請求される相談例がある"),
    ("残す物と処分する物を、どう区別しますか。作業範囲はどこまでですか", "仕分け・搬出・清掃・買取のどこまでか。処分しないよう頼んだ物の取り違えを防ぐ"),
    ("家電4品目（エアコン・テレビ・冷蔵庫/冷凍庫・洗濯機/衣類乾燥機）は、どう扱いますか", "リサイクル料金と運搬料金を、見積もりでどう扱うか。市では出せない品目"),
    ("立会いは必要ですか。当日、見積もりと違う提案があった場合はどうなりますか", "鍵の預け方と返却も確認。見積もりと異なる提案は作業前に断り、その場の支払いも拒否してよい"),
]
MITSU_NOTE = ("見積もりは複数社から取り、見積書は作業内容と費用が分かる書面で受け取ってください（国民生活センターの案内）。"
              "トラブルが起きたら、消費生活センターに相談してください。業者名は、市の公式の許可業者一覧と照らして記入してください。")


def build_mitsumori_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "見積もり比較"
    ws["A1"] = "業者見積もり比較シート（新潟版）"
    ws["A1"].font = Font(name=FONT, size=14, bold=True)
    ws["A2"] = f"{SITE}／更新 {UPDATED}／登録不要・無料・再配布可（出典明記）。薄い灰色の列は確認の目安、白い欄に記入してください。"
    ws["A2"].font = Font(name=FONT, size=9, color="6B6A66")
    ws["A3"] = "依頼する実家の所在地:"
    ws["A3"].font = Font(name=FONT, size=10)
    ws["D3"] = "記入日:"
    ws["D3"].font = Font(name=FONT, size=10)
    hr = 5
    for i, h in enumerate(["項目", "確認の目安", "業者A", "業者B", "業者C", "確認点のメモ"], 1):
        style_cell(ws.cell(hr, i, h), bold=True, fill=HEAD_FILL, valign="center")
    pos = {}
    r = hr + 1
    for name, hint, kind in MITSUMORI_ROWS:
        pos[name] = r
        style_cell(ws.cell(r, 1, name), bold=True, valign="center" if kind != "text" else "top")
        style_cell(ws.cell(r, 2, hint), size=9, color="6B6A66", fill=HINT_FILL)
        for col in range(3, 7):
            style_cell(ws.cell(r, col), size=10)
        if kind == "yen":
            for col in range(3, 6):
                ws.cell(r, col).number_format = '#,##0'
                ws.cell(r, col).alignment = Alignment(horizontal="right", vertical="center")
        if kind.startswith("list:"):
            dv = DataValidation(type="list", formula1='"' + kind[5:] + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"C{r}:E{r}")
        if kind in ("yen", "total", "name"):
            ws.row_dimensions[r].height = 22
        elif kind.startswith("list:"):
            ws.row_dimensions[r].height = 40
        else:
            ws.row_dimensions[r].height = 48
        r += 1
    tr = pos["合計（円）"]
    b, a, k = pos["基本料金（円）"], pos["追加料金の見込み額（円）"], pos["買取金額（円）"]
    for col in "CDE":
        c = ws[f"{col}{tr}"]
        c.value = f'=IF(COUNT({col}{b},{col}{a},{col}{k})=0,"",N({col}{b})+N({col}{a})-N({col}{k}))'
        c.number_format = '#,##0'
        c.font = Font(name=FONT, size=11, bold=True)
        c.alignment = Alignment(horizontal="right", vertical="center")
    ws.cell(tr, 1).fill = HEAD_FILL
    r += 1
    ws.cell(r, 1, MITSU_NOTE).font = Font(name=FONT, size=9, color="6B6A66")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 30
    for col, w in zip("ABCDEF", [26, 40, 22, 22, 22, 30]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = f"C{hr + 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4

    ws2 = wb.create_sheet("見積もりで聞く6つのこと")
    ws2["A1"] = "見積もりで聞く6つのこと"
    ws2["A1"].font = Font(name=FONT, size=14, bold=True)
    ws2["A2"] = "電話や訪問のときに、そのまま読み上げて回答を記入してください。国民生活センターの相談事例と案内に基づく項目です。"
    ws2["A2"].font = Font(name=FONT, size=9, color="6B6A66")
    for i, h in enumerate(["No.", "聞くこと", "見るポイント", "業者Aの回答", "業者Bの回答", "業者Cの回答"], 1):
        style_cell(ws2.cell(4, i, h), bold=True, fill=HEAD_FILL, valign="center")
    for n, (q, why) in enumerate(SIX, 1):
        rr = 4 + n
        style_cell(ws2.cell(rr, 1, n), align="center")
        style_cell(ws2.cell(rr, 2, q), bold=True)
        style_cell(ws2.cell(rr, 3, why), size=9, color="6B6A66", fill=HINT_FILL)
        for col in range(4, 7):
            style_cell(ws2.cell(rr, col))
        ws2.row_dimensions[rr].height = 62
    for col, w in zip("ABCDEF", [5, 38, 44, 26, 26, 26]):
        ws2.column_dimensions[col].width = w
    ws2.freeze_panes = "A5"
    ws2.page_setup.orientation = "landscape"
    ws2.page_setup.paperSize = ws2.PAPERSIZE_A4
    ws2.page_setup.fitToWidth = 1
    ws2.page_setup.fitToHeight = 0
    ws2.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(os.path.join(HERE, "mitsumori-hikaku.xlsx"))


CSS_MITSU = BASE_CSS + """
.wrap{max-width:1100px}
table.cmp td.item{width:15%;font-weight:700}
table.cmp td.hint{width:23%;font-size:12px;color:var(--sub);background:var(--soft)}
table.cmp td.v{height:34px}
table.cmp tr.tall td.v{height:62px}
table.six td.no{width:30px;text-align:center}
table.six td.q{width:30%;font-weight:700}
table.six td.why{width:26%;font-size:12px;color:var(--sub);background:var(--soft)}
table.six td.ans{height:70px}
.who{display:flex;gap:16px;flex-wrap:wrap;margin:8px 0 4px;font-size:13px}
.who span{border-bottom:1px solid var(--ink);min-width:150px;display:inline-block;height:1.4em}
.pg2{margin-top:28px}
@page{size:A4 landscape;margin:10mm}
@media print{
  body{font-size:9.5pt;line-height:1.4;color:#000}
  .wrap{max-width:none;padding:0}
  .bar{display:none}
  h1{font-size:15pt}
  h2{font-size:11pt;border-bottom:1.5px solid #000;margin-top:0}
  .pg2{page-break-before:always;margin-top:0}
  th,td{border-color:#000;padding:3px 5px}
  th{background:#fff;border-bottom:1.5px solid #000}
  td.hint,td.why{background:#fff;color:#000}
  tr{page-break-inside:avoid}
  .meta,.note,.footer{color:#000}
}
"""


def build_mitsumori_html():
    rows = []
    for name, hint, kind in MITSUMORI_ROWS:
        cls = "tall" if kind == "text" else ""
        rows.append(f'<tr class="{cls}"><td class="item">{e(name)}</td><td class="hint">{e(hint)}</td><td class="v"></td><td class="v"></td><td class="v"></td><td class="v"></td></tr>')
    six = "".join(
        f'<tr><td class="no">{n}</td><td class="q"><input type="checkbox" aria-label="聞いた"> {e(q)}</td><td class="why">{e(w)}</td><td class="ans"></td><td class="ans"></td><td class="ans"></td></tr>'
        for n, (q, w) in enumerate(SIX, 1)
    )
    doc = f"""<!doctype html>
<html lang="ja">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
{ROBOTS}
<title>業者見積もり比較シート（新潟版）</title>
<style>{CSS_MITSU}</style>
</head>
<body>
<div class="wrap">
<h1>業者見積もり比較シート（新潟版）</h1>
<p class="meta">{SITE}／更新 {UPDATED}／A4横・白黒印刷用／登録不要・無料・再配布可（出典明記）。表計算で入力する場合は同じ内容の mitsumori-hikaku.xlsx を使ってください（合計は自動計算）。</p>
<div class="bar"><button class="btn" onclick="window.print()">印刷する</button></div>
<div class="who"><div>実家の所在地: <span></span></div><div>記入日: <span></span></div></div>
<h2>見積もり比較（3社分）</h2>
<table class="cmp"><thead><tr><th>項目</th><th>確認の目安</th><th>業者A</th><th>業者B</th><th>業者C</th><th>確認点のメモ</th></tr></thead><tbody>{''.join(rows)}</tbody></table>
<h2 class="pg2">見積もりで聞く6つのこと</h2>
<table class="six"><thead><tr><th>No.</th><th>聞くこと</th><th>見るポイント</th><th>業者Aの回答</th><th>業者Bの回答</th><th>業者Cの回答</th></tr></thead><tbody>{six}</tbody></table>
<p class="footer">{e(MITSU_NOTE)}<br>出典: {SITE}「新潟の実家の片付けは自分で？業者？」（環境省・国民生活センターの公開情報に基づく）。このシートは出典を明記すれば自由に再配布できます。</p>
</div>
</body>
</html>
"""
    with open(os.path.join(HERE, "mitsumori-hikaku.html"), "w", encoding="utf-8", newline="\n") as f:
        f.write(doc)


if __name__ == "__main__":
    build_checklist_html()
    build_checklist_xlsx()
    build_mitsumori_html()
    build_mitsumori_xlsx()
    print("ok")
